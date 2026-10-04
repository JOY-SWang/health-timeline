#!/usr/bin/env python3
"""Export the canonical health timeline (timeline.json) to a human-readable Excel workbook.

Usage:
  export_excel.py --db DIR [--out FILE.xlsx] [--subject self] [--csv]

The workbook is a *view*: it is regenerated from JSON every time and should never be
edited as a source of truth. Sheets:
  Timeline    Date | Fact | Source | Analysis / Inference | Image / Evidence | Status | Event ID
              (a Subject column is added when the database tracks more than one person)
  Lab trends  one row per analyte, one column per date (only if lab results exist)
  Evidence    the archived source files
  Log         ingestion runs and every status change
  Guide       how to read the colours and columns
--csv writes timeline.csv (Timeline sheet only) for environments without openpyxl.
"""
import argparse
import csv
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from validate_record import DEFAULT_SOURCE_LABELS, resolve_db_path  # noqa: E402

ANALYSIS_PREFIX = {"ai_in_chat": "AI (chat)", "recorder": "AI (recorder)"}
INACTIVE = {"superseded", "retracted"}


# ----------------------------------------------------------------------------- row building

def sort_key(ev):
    return (ev.get("event_date") or "9999-99-99", ev.get("panel") or "", ev["event_id"])


def date_cell(ev):
    d, p = ev.get("event_date"), ev.get("date_precision")
    if p == "unknown" or not d:
        return "unknown"
    return f"≈{d}" if p == "approximate" else d


def file_evidence(ev):
    return [e for e in ev.get("evidence") or [] if e.get("kind") == "file"]


def source_cell(ev, sources):
    label = ev.get("source_label") or DEFAULT_SOURCE_LABELS.get(ev.get("source_type"), "")
    refs = []
    for e in file_evidence(ev):
        src = sources.get(e.get("source_id"), {})
        name = src.get("filename") or src.get("title") or e.get("source_id")
        refs.append(f"{name}, {e['locator']}" if e.get("locator") else name)
    return f"{label} — {'; '.join(refs)}" if refs else label


def analysis_cell(ev, by_id=None):
    lines = []
    for a in ev.get("analysis") or []:
        line = f"{ANALYSIS_PREFIX.get(a.get('by'), 'AI')}: {a.get('text', '')}"
        # an analysis keeps its original wording, but flag it when what it was based on has since changed
        stale = []
        for ref in a.get("basis") or []:
            b = (by_id or {}).get(ref, {})
            if b.get("status") == "superseded":
                stale.append(f"{ref} superseded by {b.get('superseded_by')}")
            elif b.get("status") == "retracted":
                stale.append(f"{ref} retracted")
        if stale:
            line += f" [⚠ basis changed: {'; '.join(stale)}]"
        lines.append(line)
    if ev.get("about") and ev.get("event_type") == "ai_analysis":
        lines.append(f"(about {', '.join(ev['about'])})")
    return "\n".join(lines)


def evidence_cell(ev, sources):
    parts, link = [], None
    for e in ev.get("evidence") or []:
        if e.get("kind") == "file":
            src = sources.get(e.get("source_id"), {})
            name = src.get("filename") or src.get("title") or e.get("source_id")
            parts.append(f"{name}" + (f" ({e['locator']})" if e.get("locator") else ""))
            link = link or src.get("stored_path")
        else:
            q = (e.get("quote") or "").replace("\n", " ")
            q = q if len(q) <= 80 else q[:79] + "…"
            parts.append(f"Chat ({e.get('speaker', '?')}): “{q}”")
    return "\n".join(parts), link


def status_cell(ev):
    s = ev.get("status", "")
    if s == "superseded" and ev.get("superseded_by"):
        return f"superseded → {ev['superseded_by']}"
    if s == "disputed" and ev.get("conflicts_with"):
        return f"disputed ↔ {', '.join(ev['conflicts_with'])}"
    if ev.get("supersedes"):
        return f"{s.replace('_', ' ')} (corrects {', '.join(ev['supersedes'])})"
    return s.replace("_", " ")


def timeline_rows(db, subject=None):
    sources = {s["source_id"]: s for s in db.get("sources", [])}
    names = {s["subject_id"]: s.get("display_name") or s["subject_id"] for s in db.get("subjects", [])}
    by_id = {e["event_id"]: e for e in db.get("events", [])}
    events = [e for e in db.get("events", []) if not subject or e.get("subject_id") == subject]
    multi = subject is None and len({e.get("subject_id") for e in db.get("events", [])} | set(names)) > 1
    rows = []
    for ev in sorted(events, key=sort_key):
        ev_text, link = evidence_cell(ev, sources)
        row = {
            "Date": date_cell(ev),
            "Subject": names.get(ev.get("subject_id"), ev.get("subject_id")),
            "Fact": ev.get("fact") or "—",
            "Source": source_cell(ev, sources),
            "Analysis / Inference": analysis_cell(ev, by_id),
            "Image / Evidence": ev_text,
            "Status": status_cell(ev),
            "Event ID": ev["event_id"],
            "_ev": ev,
            "_link": link,
        }
        rows.append(row)
    cols = ["Date"] + (["Subject"] if multi else []) + [
        "Fact", "Source", "Analysis / Inference", "Image / Evidence", "Status", "Event ID"]
    return cols, rows


def lab_matrix(db, subject=None):
    """{(subject, name, unit): {date: [cell text]}} for live lab results."""
    out, ref = {}, {}
    for ev in sorted(db.get("events", []), key=sort_key):
        if ev.get("event_type") != "lab_result" or ev.get("status") in INACTIVE:
            continue
        if subject and ev.get("subject_id") != subject:
            continue
        v = ev.get("values") or {}
        if not v.get("name"):
            continue
        key = (ev.get("subject_id"), v["name"], v.get("unit") or "")
        txt = f"{v.get('value')}" + (f" {v['flag']}" if v.get("flag") else "")
        if ev.get("status") == "disputed":
            txt += " (disputed)"
        elif ev.get("status") == "needs_review":
            txt += " (check)"
        out.setdefault(key, {}).setdefault(date_cell(ev), []).append(txt)
        if v.get("ref_range"):
            ref[key] = v["ref_range"]
    return out, ref


# ----------------------------------------------------------------------------- writers

def export_csv(db, out_path, subject=None):
    cols, rows = timeline_rows(db, subject)
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            w.writerow([r[c] for c in cols])
    return out_path


def export_xlsx(db, out_path, subject=None):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F4F6F")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(bottom=thin)
    analysis_font = Font(italic=True, color="5B3F8C")
    grey_font = Font(color="9AA0A6")
    grey_strike = Font(color="9AA0A6", strike=True)
    fills = {
        "disputed": PatternFill("solid", fgColor="FDE2CC"),
        "needs_review": PatternFill("solid", fgColor="FFF4C2"),
        "ai_analysis": PatternFill("solid", fgColor="F1ECFA"),
    }
    widths = {"Date": 12, "Subject": 10, "Fact": 46, "Source": 28, "Analysis / Inference": 44,
              "Image / Evidence": 32, "Status": 20, "Event ID": 11}

    def header(ws, cols, row=1):
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=row, column=i, value=c)
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")

    wb = Workbook()

    # --- Timeline
    ws = wb.active
    ws.title = "Timeline"
    cols, rows = timeline_rows(db, subject)
    header(ws, cols)
    for r_i, r in enumerate(rows, 2):
        ev = r["_ev"]
        inactive = ev.get("status") in INACTIVE
        fill = fills.get(ev.get("status")) or (fills["ai_analysis"] if ev.get("event_type") == "ai_analysis" else None)
        for c_i, c in enumerate(cols, 1):
            cell = ws.cell(row=r_i, column=c_i, value=r[c])
            cell.alignment, cell.border = wrap, border
            if c == "Analysis / Inference":
                cell.font = analysis_font if not inactive else grey_font
            elif inactive:
                cell.font = grey_strike if c == "Fact" else grey_font
            if fill is not None and not inactive:
                cell.fill = fill
            if c == "Date" and ev.get("date_note"):
                cell.comment = Comment(f"Date note: {ev['date_note']}", "health-timeline")
            if c == "Fact" and ev.get("verbatim"):
                cell.comment = Comment(f"Verbatim: {ev['verbatim']}", "health-timeline")
            if c == "Image / Evidence" and r["_link"]:
                cell.hyperlink = r["_link"]
                if not inactive:
                    cell.font = Font(color="1F5FBF", underline="single")
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 16)
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    for sheet_ws in (ws,):
        sheet_ws.page_setup.orientation = "landscape"
        sheet_ws.page_setup.paperSize = sheet_ws.PAPERSIZE_A4
        sheet_ws.sheet_properties.pageSetUpPr.fitToPage = True
        sheet_ws.page_setup.fitToWidth = 1
        sheet_ws.page_setup.fitToHeight = 0
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    # --- Lab trends
    matrix, refs = lab_matrix(db, subject)
    if matrix:
        lt = wb.create_sheet("Lab trends")
        dates = sorted({d for by_date in matrix.values() for d in by_date}, key=lambda d: d.lstrip("≈"))
        names = {s["subject_id"]: s.get("display_name") or s["subject_id"] for s in db.get("subjects", [])}
        multi = len({k[0] for k in matrix}) > 1
        lcols = (["Subject"] if multi else []) + ["Analyte", "Unit", "Ref range"] + dates
        header(lt, lcols)
        for r_i, key in enumerate(sorted(matrix, key=lambda k: (k[0], k[1].lower())), 2):
            vals = ([names.get(key[0], key[0])] if multi else []) + [key[1], key[2], refs.get(key, "")]
            vals += ["; ".join(matrix[key].get(d, [])) for d in dates]
            for c_i, v in enumerate(vals, 1):
                cell = lt.cell(row=r_i, column=c_i, value=v)
                cell.alignment, cell.border = wrap, border
        for i, c in enumerate(lcols, 1):
            lt.column_dimensions[get_column_letter(i)].width = 22 if c == "Analyte" else 13
        lt.freeze_panes = lt.cell(row=2, column=len(lcols) - len(dates) + 1)

    # --- Evidence
    es = wb.create_sheet("Evidence")
    ecols = ["Source ID", "Kind", "Title", "Filename", "Archived copy", "Document date", "Issuer",
             "Origin", "SHA-256", "Added", "Run"]
    header(es, ecols)
    for r_i, s in enumerate(db.get("sources", []), 2):
        vals = [s.get("source_id"), s.get("kind"), s.get("title"), s.get("filename"), s.get("stored_path"),
                s.get("document_date"), s.get("issuer"), s.get("origin"), (s.get("sha256") or "")[:12],
                (s.get("added_at") or "")[:16].replace("T", " "), s.get("run_id")]
        for c_i, v in enumerate(vals, 1):
            cell = es.cell(row=r_i, column=c_i, value=v)
            cell.alignment = wrap
            if c_i == 5 and v:
                cell.hyperlink = v
                cell.font = Font(color="1F5FBF", underline="single")
    for i, w in enumerate([10, 9, 30, 24, 34, 13, 18, 18, 14, 16, 10], 1):
        es.column_dimensions[get_column_letter(i)].width = w

    # --- Log
    lg = wb.create_sheet("Log")
    rcols = ["Run ID", "Run at", "Scope", "Conversation", "Last message anchor", "Events added",
             "Status changes", "Duplicates skipped", "Notes"]
    header(lg, rcols)
    row = 2
    for r in db.get("runs", []):
        ids = r.get("events_added") or []
        vals = [r.get("run_id"), (r.get("run_at") or "")[:16].replace("T", " "), r.get("scope"),
                r.get("conversation_hint"), r.get("last_message_anchor"),
                f"{len(ids)}" + (f" ({ids[0]}..{ids[-1]})" if ids else ""), r.get("status_changes"),
                r.get("duplicates_skipped"), r.get("notes")]
        for c_i, v in enumerate(vals, 1):
            lg.cell(row=row, column=c_i, value=v).alignment = wrap
        row += 1
    row += 1
    ccols = ["Changed at", "Run ID", "Event ID", "From", "To", "Reason"]
    header(lg, ccols, row=row)
    for c in db.get("changelog", []):
        row += 1
        vals = [(c.get("at") or "")[:16].replace("T", " "), c.get("run_id"), c.get("event_id"),
                c.get("from"), c.get("to"), c.get("reason")]
        for c_i, v in enumerate(vals, 1):
            lg.cell(row=row, column=c_i, value=v).alignment = wrap
    for i, w in enumerate([10, 16, 22, 26, 34, 26, 14, 12, 30], 1):
        lg.column_dimensions[get_column_letter(i)].width = w

    # --- Guide
    gd = wb.create_sheet("Guide")
    guide = [
        ("This workbook is generated from timeline.json. Edit through the health-timeline skill, not here.", None),
        ("", None),
        ("Fact", "What a source asserted, attributed to that source. A fact records that something was said or "
                 "measured - not that it is medically true."),
        ("Source", "Who or what asserted the fact (patient, doctor relayed by the user, lab report, …) plus file/page."),
        ("Analysis / Inference", "Purple italics = AI-generated. 'AI (chat)' = something the assistant said in the "
                                 "conversation; 'AI (recorder)' = a note added while recording (e.g. comparison with "
                                 "the printed reference range). Never a diagnosis; not medical advice."),
        ("Image / Evidence", "The original evidence: archived file (click to open) or a verbatim chat quote."),
        ("Status", "active · needs review (yellow: check the transcription) · disputed (orange: two sources "
                   "disagree, nothing was overwritten) · superseded (grey, struck through: replaced by a "
                   "corrected record) · retracted (grey: withdrawn)."),
        ("Date", "When it happened (clinical date), not when it was recorded. ≈ = approximate; hover for the "
                 "date note. Hover a Fact for the verbatim source text."),
        ("Purple rows", "Standalone AI analysis events (no fact of their own)."),
    ]
    for r_i, (k, v) in enumerate(guide, 1):
        gd.cell(row=r_i, column=1, value=k).font = Font(bold=bool(v))
        if v:
            gd.cell(row=r_i, column=2, value=v).alignment = wrap
    gd.column_dimensions["A"].width = 22
    gd.column_dimensions["B"].width = 100

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    tmp = out_path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, out_path)
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", help="data directory or timeline.json (default: $HEALTH_TIMELINE_DB)")
    ap.add_argument("--out", help="output path (default: timeline.xlsx / timeline.csv next to the database)")
    ap.add_argument("--subject", help="only this subject_id")
    ap.add_argument("--csv", action="store_true", help="write CSV instead of xlsx")
    args = ap.parse_args(argv)

    db_path = resolve_db_path(args.db)
    if not db_path or not os.path.exists(db_path):
        print(f"No database found at {db_path!r}", file=sys.stderr)
        return 2
    with open(db_path, "r", encoding="utf-8") as f:
        db = json.load(f)
    data_dir = os.path.dirname(db_path)
    suffix = f"_{args.subject}" if args.subject else ""
    if args.csv:
        out = export_csv(db, args.out or os.path.join(data_dir, f"timeline{suffix}.csv"), args.subject)
    else:
        try:
            out = export_xlsx(db, args.out or os.path.join(data_dir, f"timeline{suffix}.xlsx"), args.subject)
        except ImportError:
            print("openpyxl is not installed. Install it (pip install openpyxl) or use --csv.", file=sys.stderr)
            return 2
    print(f"Exported {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
