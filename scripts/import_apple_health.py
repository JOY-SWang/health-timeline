#!/usr/bin/env python3
"""Stream an Apple Health ZIP/export.xml into daily summaries and an optional timeline batch.

No network calls, archive extraction, diagnoses, or automatic database writes.
See references/apple-health.md for coverage and aggregation semantics.
"""
import argparse
from collections import Counter, defaultdict
from contextlib import contextmanager
import csv
from datetime import date, datetime
from decimal import Decimal, DecimalException
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

QUANTITY = "HKQuantityTypeIdentifier"
SLEEP = "HKCategoryTypeIdentifierSleepAnalysis"
DEFAULT_BATCH_TYPES = {QUANTITY + name for name in (
    "BodyMass", "BodyMassIndex", "Height", "RestingHeartRate", "VO2Max",
    "BloodPressureSystolic", "BloodPressureDiastolic", "BloodGlucose",
)}
ADDITIVE_TYPES = {QUANTITY + name for name in (
    "StepCount", "DistanceWalkingRunning", "DistanceCycling", "DistanceSwimming",
    "BasalEnergyBurned", "ActiveEnergyBurned", "FlightsClimbed", "AppleExerciseTime",
    "AppleStandTime", "TimeInDaylight",
)}
SLEEP_STAGES = {"HKCategoryValueSleepAnalysis" + stage: stage for stage in (
    "InBed", "Awake", "Asleep", "AsleepUnspecified", "AsleepCore", "AsleepDeep", "AsleepREM",
)}
CSV_FIELDS = ["date", "type", "source", "unit", "stage", "count", "minimum", "maximum",
              "sample_mean", "recorded_sum", "duration_hours"]


class EntityGuard:
    """Apple's internal element DTD is valid; custom entities and non-UTF-8 XML are not."""
    def __init__(self, stream):
        self.stream = stream
        self.tail = b""

    def read(self, size=-1):
        chunk = self.stream.read(size)
        scanned = self.tail + chunk
        if b"<!ENTITY" in scanned.upper():
            raise ValueError("XML entity declarations are not supported")
        if b"\0" in chunk:
            raise ValueError("Only UTF-8 Apple Health XML is supported")
        self.tail = scanned[-16:]
        return chunk


@contextmanager
def export_stream(path):
    """Open one XML member in place. ZIP paths are never written to the filesystem."""
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            entries = [m for m in archive.infolist() if not m.is_dir()]
            candidates = [m for m in entries if m.filename.rsplit("/", 1)[-1] == "export.xml"]
            if len(candidates) != 1:
                raise ValueError("ZIP must contain exactly one export.xml; found %d" % len(candidates))
            member = candidates[0]
            members = [{"name": m.filename, "bytes": m.file_size, "processed": m is member}
                       for m in entries]
            with archive.open(member) as stream:
                yield EntityGuard(stream), member.filename, members
    else:
        if path.suffix.lower() != ".xml":
            raise ValueError("Input must be an Apple Health ZIP or export.xml")
        with path.open("rb") as stream:
            yield EntityGuard(stream), path.name, [{"name": path.name, "bytes": path.stat().st_size,
                                                   "processed": True}]


def parse_timestamp(value):
    # Normalise Apple's space before +/-HHMM for Python 3.8's fromisoformat.
    value = re.sub(r"\s*([+-]\d{2})(\d{2})$", r"\1:\2", value or "")
    result = datetime.fromisoformat(value)
    if result.tzinfo is None:
        raise ValueError("timestamp has no UTC offset")
    return result


def type_set(text):
    if text is None:
        return None
    names = [name.strip() for name in text.split(",") if name.strip()]
    if not names:
        raise argparse.ArgumentTypeError("specify at least one type")
    return {name if name.startswith("HK") else SLEEP if name == "SleepAnalysis" else QUANTITY + name
            for name in names}


def decimal_text(value):
    # Summary statistics are derived, and means/durations are rounded to six decimal places.
    if value == 0:
        return "0"
    return format(value, "f").rstrip("0").rstrip(".") if value.as_tuple().exponent < 0 else format(value, "f")


def rounded(value):
    return decimal_text(value.quantize(Decimal("0.000001")))


class QuantityStats:
    def __init__(self):
        self.count = 0
        self.total = Decimal(0)
        self.minimum = None
        self.maximum = None

    def add(self, value):
        self.count += 1
        self.total += value
        self.minimum = value if self.minimum is None else min(self.minimum, value)
        self.maximum = value if self.maximum is None else max(self.maximum, value)

    def row(self, key):
        day, kind, source, unit = key
        return dict(date=day, type=kind, source=source, unit=unit, stage="", count=self.count,
                    minimum=decimal_text(self.minimum), maximum=decimal_text(self.maximum),
                    sample_mean=rounded(self.total / self.count),
                    recorded_sum=decimal_text(self.total) if kind in ADDITIVE_TYPES else "",
                    duration_hours="")


def union_seconds(intervals):
    seconds = 0
    left = right = None
    for start, end in sorted(intervals):
        if right is not None and start <= right:
            right = max(right, end)
        else:
            if right is not None:
                seconds += (right - left).total_seconds()
            left, right = start, end
    if right is not None:
        seconds += (right - left).total_seconds()
    return Decimal(str(seconds))


def measurement_event(attrs, ordinal, member, subject):
    kind = attrs["type"]
    name = kind[len(QUANTITY):]
    source = attrs.get("sourceName") or "Unknown source"
    value, unit = attrs["value"], attrs.get("unit", "")
    start, end = attrs["startDate"], attrs["endDate"]
    # The legacy fact fingerprint strips decimal points, punctuation and offset signs.
    # Preserve an exact opt-in identity without changing fingerprints of existing events.
    identity = json.dumps([kind, source, value, unit, start, end], ensure_ascii=False, separators=(",", ":"))
    measurement_id = "sha256:" + hashlib.sha256(identity.encode("utf-8")).hexdigest()
    fact = "Apple Health export (%s): %s %s %s; start=%s; end=%s" % (source, name, value, unit, start, end)
    missing = not unit or not attrs.get("sourceName")
    return {
        "temp_id": "A%d" % ordinal, "subject_id": subject,
        "event_date": start[:10], "date_precision": "day",
        "date_note": "Local start date and UTC offset as recorded in the Apple Health export.",
        "event_type": "vital_sign", "source_type": "device_measurement",
        "source_label": "Apple Health / " + source, "fact": fact,
        "verbatim": json.dumps(attrs, ensure_ascii=False, sort_keys=True),
        "values": {"name": name, "value": value, "unit": unit},
        "measurement_id": measurement_id,
        "evidence": [{"kind": "file", "source_id": "S1",
                      "locator": "%s; Record #%d (1-based XML document order); type=%s; startDate=%s" %
                                 (member, ordinal, kind, start)}],
        "status": "needs_review" if missing else "active",
        "confidence": "medium" if missing else "high",
        "confidence_note": "Source name or unit missing in export." if missing else "",
    }


def read_export(path, since=None, until=None, types=None, batch=False,
                batch_types=None, max_events=10000, subject="self", anchor=None, progress_every=250000):
    quantity = defaultdict(QuantityStats)
    sleep = defaultdict(list)
    type_counts, other_counts, warnings = Counter(), Counter(), Counter()
    coverage = dict(record_count=0, processed_records=0, filtered_records=0,
                    invalid_records=0, unsupported_records=0)
    events = []
    batch_types = DEFAULT_BATCH_TYPES if batch_types is None else batch_types
    min_day = max_day = None
    export_date = None
    with export_stream(path) as (stream, member, members):
        depth = 0
        for action, element in ET.iterparse(stream, events=("start", "end")):
            if action == "start":
                depth += 1
                if depth == 1:
                    root = element
                    if root.tag != "HealthData":
                        raise ValueError("Expected Apple Health HealthData XML root (not a CDA clinical document)")
                continue
            if element.tag == "ExportDate":
                export_date = element.get("value")
            elif element.tag == "Record":
                coverage["record_count"] += 1
                ordinal = coverage["record_count"]
                attrs = element.attrib
                kind = attrs.get("type", "unknown")
                type_counts[kind] += 1
                if progress_every and ordinal % progress_every == 0:
                    print("Read %d records" % ordinal, file=sys.stderr)
                try:
                    start = parse_timestamp(attrs.get("startDate"))
                    end = parse_timestamp(attrs.get("endDate"))
                    if end < start:
                        raise ValueError("end before start")
                    day = (end if kind == SLEEP else start).date().isoformat()
                except (ValueError, TypeError, OverflowError):
                    coverage["invalid_records"] += 1
                    warnings["invalid_timestamp"] += 1
                else:
                    min_day = min(min_day or day, day)
                    max_day = max(max_day or day, day)
                    if (since and day < since) or (until and day > until) or (types is not None and kind not in types):
                        coverage["filtered_records"] += 1
                    elif kind.startswith(QUANTITY):
                        try:
                            raw = attrs.get("value", "")
                            if len(raw) > 128:
                                raise ValueError("numeric representation too long")
                            value = Decimal(raw)
                            if not value.is_finite():
                                raise ValueError("invalid numeric value")
                            digits, exponent = value.as_tuple().digits, value.as_tuple().exponent
                            if len(digits) > 28 or not -30 <= exponent <= 18 or value.copy_abs() > Decimal("1e18"):
                                raise ValueError("numeric representation exceeds working range")
                        except (DecimalException, ValueError):
                            coverage["invalid_records"] += 1
                            warnings["invalid_numeric_value"] += 1
                        else:
                            source, unit = attrs.get("sourceName") or "Unknown source", attrs.get("unit", "")
                            quantity[(day, kind, source, unit)].add(value)
                            coverage["processed_records"] += 1
                            if not attrs.get("sourceName") or not unit:
                                warnings["missing_source_or_unit"] += 1
                            if batch and kind in batch_types:
                                if len(events) >= max_events:
                                    raise ValueError("Timeline batch event limit exceeded; narrow --since/--until/--batch-types")
                                events.append(measurement_event(dict(attrs), ordinal, member, subject))
                    elif kind == SLEEP and attrs.get("value") in SLEEP_STAGES:
                        stage = SLEEP_STAGES[attrs["value"]]
                        source = attrs.get("sourceName") or "Unknown source"
                        if not attrs.get("sourceName"):
                            warnings["missing_sleep_source"] += 1
                        sleep[(day, source, stage)].append((start, end))
                        if stage.startswith("Asleep"):
                            sleep[(day, source, "asleep_total")].append((start, end))
                        coverage["processed_records"] += 1
                    else:
                        coverage["unsupported_records"] += 1
                element.clear()
            elif depth == 2:
                other_counts[element.tag] += 1
            if depth == 2:
                root.clear()  # Remove completed siblings; clearing Record alone leaks millions of nodes.
            depth -= 1

    daily = [stats.row(key) for key, stats in sorted(quantity.items())]
    for (day, source, stage), intervals in sorted(sleep.items()):
        daily.append(dict(date=day, type=SLEEP, source=source, unit="h", stage=stage,
                          count=len(intervals), minimum="", maximum="", sample_mean="", recorded_sum="",
                          duration_hours=rounded(union_seconds(intervals) / 3600)))
    daily.sort(key=lambda r: (r["date"], r["type"], r["source"], r["unit"], r["stage"]))
    digest = hashlib.sha256()
    with path.open("rb") as original:
        for chunk in iter(lambda: original.read(1 << 20), b""):
            digest.update(chunk)
    coverage.update(record_types=dict(sorted(type_counts.items())), other_elements=dict(sorted(other_counts.items())),
                    first_local_date=min_day, last_local_date=max_day)
    report = {
        "format": "apple-health-summary/1", "input": {"path": str(path), "sha256": digest.hexdigest(),
            "xml_member": member, "members": members, "export_date": export_date},
        "selection": {"since": since, "until": until, "types": sorted(types) if types is not None else None},
        "coverage": coverage, "warnings": dict(sorted(warnings.items())), "daily": daily,
        "semantics": {
            "quantity_day": "Record startDate in its recorded UTC offset; intervals are not split at midnight.",
            "quantity_statistics": "Raw exported samples, without de-duplication or source priority; sample_mean is not time weighted.",
            "recorded_sum": "Sum of raw additive samples per source and unit, not an Apple Health daily total; overlapping samples can double count.",
            "sleep_day": "Each entire sleep interval is assigned to its local end date (wake date).",
            "sleep_duration": "Union of overlapping intervals per source and stage; asleep_total unions asleep stages only, excluding InBed and Awake.",
            "rounding": "Derived means and durations rounded to six decimal places; original batch values unchanged.",
            "coverage": "Only export.xml quantity Records and recognised sleep Records processed; other categories, workouts, activity summaries, CDA, ECG and GPX are not imported.",
        },
        "batch_events": len(events),
    }
    payload = None
    if batch:
        payload = {
            "run": {"scope": "Apple Health export; since=%s; until=%s" % (since or "all", until or "all"),
                    "conversation_hint": "Apple Health import",
                    "last_message_anchor": anchor or "file: %s; sha256=%s" % (path.name, digest.hexdigest())},
            "sources": [{"temp_id": "S1", "kind": "other", "title": "Apple Health export",
                         "file_path": str(path), "filename": path.name, "issuer": "Apple Health",
                         "origin": "user-provided local export", "notes": "Member %s; selected raw quantity measurements only." % member}],
            "events": events,
        }
    return report, payload


def export_summary_xlsx(path, report):
    """A separate derived-data workbook, never the canonical timeline or its fact layer."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)

    def sheet(title, fields, rows, numeric=()):
        ws = workbook.create_sheet(title)
        ws.freeze_panes = "A2"
        for index, field in enumerate(fields, 1):
            widths = {"type": 52, "source": 32, "metric": 55, "value": 100}
            ws.column_dimensions[get_column_letter(index)].width = widths.get(field, 20)
        headers = []
        for field in fields:
            cell = WriteOnlyCell(ws, value=field)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F4F6F")
            headers.append(cell)
        ws.append(headers)
        count = 0
        for row in rows:
            cells = []
            for field in fields:
                value = row.get(field, "")
                if field in numeric and value != "":
                    value = Decimal(str(value))
                cell = WriteOnlyCell(ws, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=field in ("type", "source", "value"))
                if field in numeric:
                    cell.number_format = "0.######"
                cells.append(cell)
            ws.append(cells)
            count += 1
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(fields)), count + 1)

    fields = ["date", "type", "source", "unit", "count", "minimum", "maximum", "sample_mean", "recorded_sum"]
    sheet("Daily quantities", fields, (r for r in report["daily"] if r["type"] != SLEEP),
          numeric=("count", "minimum", "maximum", "sample_mean", "recorded_sum"))
    sheet("Sleep", ["date", "source", "stage", "count", "duration_hours"],
          (r for r in report["daily"] if r["type"] == SLEEP), numeric=("count", "duration_hours"))
    coverage = [{"metric": key, "value": value} for key, value in report["coverage"].items()
                if not isinstance(value, dict)]
    coverage += [{"metric": "batch_events", "value": report["batch_events"]},
                 {"metric": "input_sha256", "value": report["input"]["sha256"]}]
    coverage += [{"metric": "warning: " + key, "value": value} for key, value in report["warnings"].items()]
    coverage += [{"metric": "Record type: " + key, "value": value}
                 for key, value in report["coverage"]["record_types"].items()]
    coverage += [{"metric": "other element: " + key, "value": value}
                 for key, value in report["coverage"]["other_elements"].items()]
    coverage += [{"metric": "member: " + m["name"], "value": "processed" if m["processed"] else "not processed"}
                 for m in report["input"]["members"]]
    sheet("Coverage", ["metric", "value"], coverage)
    guide = [{"metric": "Data", "value": "Calculated summaries of exported samples; original measurements remain in the export and optional batch."}]
    guide += [{"metric": key, "value": value} for key, value in report["semantics"].items()]
    guide += [{"metric": "Excel precision", "value": "Excel numeric cells have limited precision; JSON and original batch strings retain the authoritative decimal representations."}]
    sheet("Guide", ["metric", "value"], guide)
    workbook.save(path)


def write_outputs(out, report, batch, xlsx=False):
    if out.exists():
        raise ValueError("Output directory already exists; choose a new --out to preserve previous exports")
    out.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".apple-health-", dir=str(out.parent)))
    try:
        def write_json(name, data):
            with (staging / name).open("w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2, allow_nan=False)
                f.write("\n")
        write_json("apple-health-summary.json", report)
        with (staging / "apple-health-daily.csv").open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
            writer.writeheader()
            for row in report["daily"]:
                safe = dict(row)
                for key in ("source", "type", "unit", "stage"):
                    if str(safe[key]).lstrip().startswith(("=", "+", "-", "@")):
                        safe[key] = "'" + safe[key]
                writer.writerow(safe)
        if batch is not None:
            write_json("apple-health-batch.json", batch)
        if xlsx:
            export_summary_xlsx(staging / "apple-health-summary.xlsx", report)
        # All output files become visible together only after the entire parse succeeds.
        staging.rename(out)
    finally:
        if staging.exists():
            shutil.rmtree(staging)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Apple Health export.zip or export.xml")
    parser.add_argument("--out", required=True, type=Path, help="new local output directory")
    parser.add_argument("--since", type=date.fromisoformat, help="inclusive local YYYY-MM-DD")
    parser.add_argument("--until", type=date.fromisoformat, help="inclusive local YYYY-MM-DD")
    parser.add_argument("--types", type=type_set, help="comma-separated types for summaries (default: all quantities and sleep)")
    parser.add_argument("--batch", action="store_true", help="also write a batch of selected original measurements; never append automatically")
    parser.add_argument("--xlsx", action="store_true", help="also export a separate daily/sleep summary workbook (requires openpyxl)")
    parser.add_argument("--batch-types", type=type_set, help="comma-separated raw quantity types (default: conservative low-frequency set)")
    parser.add_argument("--max-events", type=int, default=10000)
    parser.add_argument("--subject", default="self", help="existing timeline subject ID")
    parser.add_argument("--anchor", help="verbatim last processed chat message, for the timeline checkpoint")
    parser.add_argument("--progress-every", type=int, default=250000, help="progress interval in records; 0 disables")
    args = parser.parse_args(argv)
    try:
        path, out = args.input.expanduser().resolve(), args.out.expanduser().resolve()
        if not path.is_file():
            raise ValueError("Input file does not exist")
        if out.exists():
            raise ValueError("Output directory already exists; choose a new --out")
        if args.since and args.until and args.since > args.until:
            raise ValueError("--since must not be after --until")
        if args.max_events < 1 or args.progress_every < 0:
            raise ValueError("--max-events must be positive and --progress-every nonnegative")
        if not re.fullmatch(r"[a-z0-9_-]+", args.subject):
            raise ValueError("--subject must be a timeline subject ID")
        if args.batch_types and any(not t.startswith(QUANTITY) for t in args.batch_types):
            raise ValueError("--batch-types supports original quantity measurements only")
        if args.xlsx:
            try:
                import openpyxl  # noqa: F401 -- check optional dependency before a potentially large parse
            except ImportError:
                raise ValueError("--xlsx requires openpyxl; install the existing requirements or omit --xlsx for CSV/JSON")
        report, batch = read_export(path, since=args.since.isoformat() if args.since else None,
                                   until=args.until.isoformat() if args.until else None, types=args.types,
                                   batch=args.batch, batch_types=args.batch_types, max_events=args.max_events,
                                   subject=args.subject, anchor=args.anchor, progress_every=args.progress_every)
        write_outputs(out, report, batch, xlsx=args.xlsx)
        c = report["coverage"]
        print("Read %d Records; processed %d; filtered %d; invalid %d; unsupported %d." %
              (c["record_count"], c["processed_records"], c["filtered_records"], c["invalid_records"], c["unsupported_records"]))
        print("Wrote %d daily rows and %d batch events to %s" % (len(report["daily"]), report["batch_events"], out))
        return 0
    except (OSError, ValueError, DecimalException, ET.ParseError, zipfile.BadZipFile, RuntimeError) as error:
        message = "Invalid XML: %s" % error if isinstance(error, ET.ParseError) else str(error)
        print("Apple Health import failed: %s" % message, file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
