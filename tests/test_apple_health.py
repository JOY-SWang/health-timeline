"""Synthetic exports only: never copy patient exports into this suite."""
import csv
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.sax.saxutils import quoteattr

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "import_apple_health.py"
PREFIX = "HKQuantityTypeIdentifier"
SLEEP = "HKCategoryTypeIdentifierSleepAnalysis"


def record(kind="BodyMass", value="70", unit="kg", source="Synthetic Watch",
           start="2026-01-01 08:00:00 -0500", end=None, **extra):
    attrs = dict(type=kind if kind.startswith("HK") else PREFIX + kind,
                 value=value, unit=unit, sourceName=source, sourceVersion="1",
                 creationDate=start, startDate=start, endDate=end or start)
    attrs.update(extra)
    return "<Record " + " ".join(k + "=" + quoteattr(v) for k, v in attrs.items()) + "/>"


def xml(records, extra=""):
    return ('<?xml version="1.0" encoding="UTF-8"?>\n'
            '<!DOCTYPE HealthData [<!ELEMENT HealthData ANY>]>\n'
            '<HealthData locale="en_US"><ExportDate value="2026-01-04 10:00:00 -0500"/>'
            '<Me HKCharacteristicTypeIdentifierDateOfBirth="1900-01-01"/>'
            + "".join(records) + extra + "</HealthData>")


class AppleHealthTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.out = self.dir / "result"

    def tearDown(self):
        self.tmp.cleanup()

    def export(self, records, extra="", zipped=True):
        path = self.dir / ("export.zip" if zipped else "export.xml")
        payload = xml(records, extra)
        if zipped:
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
                z.writestr("apple_health_export/export.xml", payload)
                z.writestr("apple_health_export/export_cda.xml", "<not-a-second-input/>")
                z.writestr("apple_health_export/workout-routes/route.gpx", "PRIVATE-GPS")
        else:
            path.write_text(payload, encoding="utf-8")
        return path

    def run_import(self, path, *args, success=True):
        proc = subprocess.run([sys.executable, str(SCRIPT), str(path), "--out", str(self.out), *args],
                              capture_output=True, text=True)
        if success:
            self.assertEqual(proc.returncode, 0, proc.stderr)
            return json.loads((self.out / "apple-health-summary.json").read_text())
        self.assertNotEqual(proc.returncode, 0)
        return proc

    def test_zip_inventory_and_source_unit_separation(self):
        path = self.export([
            record("HeartRate", "60", "count/min"),
            record("HeartRate", "80", "count/min", start="2026-01-01 12:00:00 -0500"),
            record("HeartRate", "90", "count/min", source="Synthetic Phone"),
            record("BodyMass", "150", "lb"), record("BodyMass", "70", "kg"),
        ], '<Workout duration="20"/><ActivitySummary dateComponents="2026-01-01"/>')
        report = self.run_import(path)
        self.assertEqual(report["coverage"]["record_count"], 5)
        self.assertEqual(report["coverage"]["other_elements"]["Workout"], 1)
        self.assertEqual(report["coverage"]["other_elements"]["ActivitySummary"], 1)
        self.assertEqual(len(report["input"]["members"]), 3)
        self.assertEqual(sum(m["processed"] for m in report["input"]["members"]), 1)
        rows = report["daily"]
        self.assertEqual(len(rows), 4)
        watch = next(r for r in rows if r["type"] == PREFIX + "HeartRate" and r["source"] == "Synthetic Watch")
        self.assertEqual((watch["count"], watch["minimum"], watch["maximum"], watch["sample_mean"]),
                         (2, "60", "80", "70"))
        self.assertEqual(watch["recorded_sum"], "")
        self.assertFalse(any(p.suffix in (".xml", ".gpx") for p in self.out.rglob("*")))
        self.assertNotIn("1900-01-01", (self.out / "apple-health-summary.json").read_text())

    def test_additive_values_are_separate_recorded_sums(self):
        report = self.run_import(self.export([
            record("StepCount", "100", "count"), record("StepCount", "200", "count"),
            record("StepCount", "50", "count", source="Synthetic Phone"),
        ]))
        self.assertEqual(sorted(r["recorded_sum"] for r in report["daily"]), ["300", "50"])

    def test_sleep_union_excludes_in_bed_and_awake_and_keeps_sources(self):
        def sleep(stage, start, end, source="Synthetic Watch"):
            return record(SLEEP, "HKCategoryValueSleepAnalysis" + stage, "", source,
                          "2026-01-01 " + start + " -0500", "2026-01-02 " + end + " -0500")
        report = self.run_import(self.export([
            sleep("AsleepCore", "23:00:00", "01:00:00"),
            sleep("AsleepCore", "23:00:00", "01:00:00"),
            record(SLEEP, "HKCategoryValueSleepAnalysisAsleepDeep", "",
                   start="2026-01-02 00:30:00 -0500", end="2026-01-02 02:00:00 -0500"),
            sleep("InBed", "22:00:00", "03:00:00"),
            record(SLEEP, "HKCategoryValueSleepAnalysisAwake", "",
                   start="2026-01-02 02:00:00 -0500", end="2026-01-02 03:00:00 -0500"),
            sleep("AsleepUnspecified", "23:00:00", "02:00:00", "Synthetic Phone"),
        ]))
        totals = [r for r in report["daily"] if r["stage"] == "asleep_total"]
        self.assertEqual(len(totals), 2)
        self.assertTrue(all(r["date"] == "2026-01-02" and r["duration_hours"] == "3" for r in totals))
        self.assertEqual(next(r["duration_hours"] for r in report["daily"] if r["stage"] == "InBed"), "5")

    def test_sleep_elapsed_time_respects_offset_change(self):
        report = self.run_import(self.export([record(
            SLEEP, "HKCategoryValueSleepAnalysisAsleepUnspecified", "",
            start="2025-11-02 00:30:00 -0400", end="2025-11-02 02:30:00 -0500")]))
        self.assertEqual(next(r["duration_hours"] for r in report["daily"] if r["stage"] == "asleep_total"), "3")

    def test_date_and_type_filters_use_record_local_date(self):
        report = self.run_import(self.export([
            record(start="2026-01-01 23:30:00 -0500"),
            record(start="2026-01-02 00:30:00 +0800"),
            record("HeartRate", "80", "count/min", start="2026-01-02 12:00:00 +0800"),
            record(start="2026-01-03 00:30:00 +0800"),
        ], zipped=False), "--since", "2026-01-02", "--until", "2026-01-02", "--types", "BodyMass")
        self.assertEqual(report["coverage"]["record_count"], 4)
        self.assertEqual(report["coverage"]["filtered_records"], 3)
        self.assertEqual(len(report["daily"]), 1)
        self.assertEqual(report["daily"][0]["count"], 1)

    def test_batch_preserves_same_day_samples_and_is_repeatable(self):
        path = self.export([record(), record(start="2026-01-01 12:00:00 -0500"),
                            record("HeartRate", "60", "count/min")])
        self.run_import(path, "--batch")
        batch_path = self.out / "apple-health-batch.json"
        batch = json.loads(batch_path.read_text())
        self.assertEqual(len(batch["events"]), 2)
        self.assertTrue(Path(batch["sources"][0]["file_path"]).samefile(path))
        self.assertEqual(batch["events"][0]["values"]["value"], "70")
        self.assertNotEqual(batch["events"][0]["fact"], batch["events"][1]["fact"])
        self.assertIn("Record #1", batch["events"][0]["evidence"][0]["locator"])
        db = self.dir / "db"
        for args in (["--init"], ["--batch", str(batch_path), "--no-export"],
                     ["--batch", str(batch_path), "--no-export"]):
            p = subprocess.run([sys.executable, str(ROOT / "scripts/append_event.py"), "--db", str(db), *args],
                               capture_output=True, text=True)
            self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        saved = json.loads((db / "timeline.json").read_text())
        self.assertEqual(len(saved["events"]), 2)
        self.assertEqual(len(saved["sources"]), 1)
        self.assertTrue((db / saved["sources"][0]["stored_path"]).exists())

    def test_invalid_values_and_unsupported_categories_are_accounted_for(self):
        report = self.run_import(self.export([
            record(value="NaN"), record(value="Infinity"), record(start="not-a-date"),
            record("HKCategoryTypeIdentifierMindfulSession", "1", ""), record(),
        ]))
        coverage = report["coverage"]
        self.assertEqual((coverage["record_count"], coverage["processed_records"],
                          coverage["invalid_records"], coverage["unsupported_records"]), (5, 1, 3, 1))

    def test_raw_identity_keeps_decimal_source_and_offset_distinctions(self):
        self.run_import(self.export([
            record(value="7.0"), record(value="70"),
            record(value="70", start="2026-01-01 08:00:00 +0500"),
            record(value="70", source="Synthetic-Watch"),
        ]), "--batch")
        db = self.dir / "db"
        for args in (["--init"], ["--batch", str(self.out / "apple-health-batch.json"), "--no-export"]):
            p = subprocess.run([sys.executable, str(ROOT / "scripts/append_event.py"), "--db", str(db), *args],
                               capture_output=True, text=True)
            self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        saved = json.loads((db / "timeline.json").read_text())
        self.assertEqual(len(saved["events"]), 4)

    def test_extreme_numeric_exponents_are_counted_without_large_output(self):
        report = self.run_import(self.export([
            record(value="1e-1000000"), record(value="1e999999999"), record(),
        ]))
        self.assertEqual(report["coverage"]["invalid_records"], 2)
        self.assertEqual(report["coverage"]["processed_records"], 1)
        self.assertLess((self.out / "apple-health-summary.json").stat().st_size, 20000)

    def test_invalid_xml_does_not_leave_partial_outputs(self):
        path = self.dir / "export.xml"
        path.write_text("<HealthData>" + record() + "<broken>")
        p = self.run_import(path, success=False)
        self.assertIn("XML", p.stderr)
        self.assertFalse(self.out.exists())

    def test_entities_are_rejected_but_apple_doctype_is_accepted(self):
        path = self.dir / "export.xml"
        path.write_text('<!DOCTYPE HealthData [<!ENTITY secret SYSTEM "file:///etc/passwd">]>'
                        '<HealthData>&secret;</HealthData>')
        p = self.run_import(path, success=False)
        self.assertIn("entity", p.stderr.lower())
        self.assertFalse(self.out.exists())

    def test_batch_limit_fails_without_partial_artifacts(self):
        p = self.run_import(self.export([record(), record()]), "--batch", "--max-events", "1", success=False)
        self.assertIn("limit", p.stderr)
        self.assertFalse(self.out.exists())

    def test_existing_outputs_are_not_overwritten(self):
        path = self.export([record()])
        self.out.mkdir()
        sentinel = self.out / "keep.txt"
        sentinel.write_text("keep")
        p = self.run_import(path, success=False)
        self.assertIn("exists", p.stderr)
        self.assertEqual(sentinel.read_text(), "keep")

    def test_csv_text_cannot_inject_spreadsheet_formulas(self):
        self.run_import(self.export([record(source='=HYPERLINK("https://invalid.test")')]))
        with (self.out / "apple-health-daily.csv").open(newline="") as f:
            row = next(csv.DictReader(f))
        self.assertTrue(row["source"].startswith("'="))

    def test_optional_summary_excel_includes_sleep_without_creating_facts(self):
        import openpyxl
        source = '=HYPERLINK("https://invalid.test")'
        self.run_import(self.export([
            record(source=source),
            record(SLEEP, "HKCategoryValueSleepAnalysisAsleepCore", "", source=source,
                   start="2026-01-01 23:00:00 -0500", end="2026-01-02 01:00:00 -0500"),
        ]), "--xlsx")
        wb = openpyxl.load_workbook(self.out / "apple-health-summary.xlsx")
        self.assertIn("Daily quantities", wb.sheetnames)
        self.assertIn("Sleep", wb.sheetnames)
        self.assertIn("Coverage", wb.sheetnames)
        sleep = list(wb["Sleep"].values)
        self.assertEqual(len(sleep), 3)
        self.assertTrue(all(row[4] == 2 for row in sleep[1:]))
        self.assertEqual(wb["Sleep"].cell(2, 2).data_type, "s")
        self.assertFalse((self.out / "apple-health-batch.json").exists())

    def test_sleep_missing_source_is_reported(self):
        report = self.run_import(self.export([record(
            SLEEP, "HKCategoryValueSleepAnalysisAsleepCore", "", source="",
            start="2026-01-01 23:00:00 -0500", end="2026-01-02 01:00:00 -0500")]))
        self.assertEqual(report["warnings"]["missing_sleep_source"], 1)

    def test_device_measurements_have_their_own_excel_trends(self):
        sys.path.insert(0, str(ROOT / "scripts"))
        import openpyxl
        self.run_import(self.export([record()]), "--batch")
        db = self.dir / "db"
        for args in (["--init"], ["--batch", str(self.out / "apple-health-batch.json")]):
            p = subprocess.run([sys.executable, str(ROOT / "scripts/append_event.py"), "--db", str(db), *args],
                               capture_output=True, text=True)
            self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        wb = openpyxl.load_workbook(db / "timeline.xlsx")
        self.assertIn("Device trends", wb.sheetnames)
        self.assertNotIn("Lab trends", wb.sheetnames)
        self.assertIn("BodyMass", " ".join(str(c.value) for row in wb["Device trends"] for c in row))


if __name__ == "__main__":
    unittest.main()
